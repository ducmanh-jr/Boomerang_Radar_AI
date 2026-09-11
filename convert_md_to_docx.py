import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout_box(doc, text, title=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=120, bottom=120, left=200, right=200)
    
    # Left border navy thick
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="0A2540"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(10, 37, 64)
        
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10.5)
    run_text.font.color.rgb = RGBColor(45, 55, 72)
    
    # Spacing after table
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(6)

def parse_inline_formatting(p, text):
    # Parses **bold**, *italic*, `code`, and $math$ inside line text
    pattern = r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$)'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(197, 48, 48)
            set_cell_background_run = True
        elif token.startswith('$') and token.endswith('$'):
            run = p.add_run(token[1:-1])
            run.italic = True
            run.font.name = "Cambria Math"
            run.font.color.rgb = RGBColor(43, 108, 176)
        else:
            run = p.add_run(token)
        run.font.name = run.font.name or "Calibri"
        run.font.size = run.font.size or Pt(11)

def convert_md_to_docx(md_path, docx_path, images_dir):
    doc = Document()
    
    # Page Margins 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_data = []
    
    in_code_block = False
    code_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Code block handling
        if line.startswith('```'):
            if in_code_block:
                # End of code block
                code_text = '\n'.join(code_lines)
                add_callout_box(doc, code_text, title="CODE / PAYLOAD SPECIFICATION")
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table handling
        if '|' in line and not line.startswith('!['):
            parts = [cell.strip() for cell in line.split('|')[1:-1]]
            if len(parts) > 0:
                # Check if separator line
                if all(re.match(r'^:?-+:?$', p) for p in parts):
                    i += 1
                    continue
                in_table = True
                table_data.append(parts)
                i += 1
                continue
        else:
            if in_table and len(table_data) > 0:
                # Process collected table
                cols = max(len(row) for row in table_data)
                rows = len(table_data)
                tbl = doc.add_table(rows=rows, cols=cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                for r_idx, row_vals in enumerate(table_data):
                    for c_idx, val in enumerate(row_vals):
                        if c_idx < cols:
                            cell = tbl.cell(r_idx, c_idx)
                            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            
                            if r_idx == 0:
                                set_cell_background(cell, "0A2540")
                                run = p.add_run(val)
                                run.bold = True
                                run.font.name = "Calibri"
                                run.font.size = Pt(10)
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                fill = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                                set_cell_background(cell, fill)
                                parse_inline_formatting(p, val)
                                
                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_before = Pt(0)
                p_sp.paragraph_format.space_after = Pt(6)
                
                table_data = []
                in_table = False

        # Blank lines
        if not line.strip():
            i += 1
            continue

        # Horizontal Rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="CBD5E0"/></w:pBdr>')
            p._element.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Images
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line.strip())
        if img_match:
            alt_text, img_rel_path = img_match.groups()
            img_abs_path = os.path.normpath(os.path.join(r"c:\Users\Admin\ducmanh\DMC_kha-nang-khach-hang-quay-lai", img_rel_path))
            if os.path.exists(img_abs_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(img_abs_path, width=Inches(6.2))
                
                # Check for caption next line
                if i + 1 < len(lines) and lines[i+1].strip().startswith('*Hình'):
                    caption_line = lines[i+1].strip()
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(0)
                    p_cap.paragraph_format.space_after = Pt(12)
                    run_cap = p_cap.add_run(caption_line.strip('*'))
                    run_cap.italic = True
                    run_cap.font.name = "Calibri"
                    run_cap.font.size = Pt(9.5)
                    run_cap.font.color.rgb = RGBColor(113, 128, 150)
                    i += 1
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(line[2:])
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(10, 37, 64) # Navy
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[3:])
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0, 128, 128) # Teal
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[4:])
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(43, 108, 176) # Blue
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[5:])
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(45, 55, 72)
        elif line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_formatting(p, line[2:])
        elif re.match(r'^\d+\.\s', line):
            content = re.sub(r'^\d+\.\s', '', line)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_formatting(p, content)
        elif line.startswith('$$') and line.endswith('$$'):
            # Display Math Block
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:-2])
            run.italic = True
            run.font.name = "Cambria Math"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(43, 108, 176)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_formatting(p, line)
            
        i += 1
        
    doc.save(docx_path)
    print(f"Successfully converted {md_path} -> {docx_path}")

if __name__ == '__main__':
    md = r"c:\Users\Admin\ducmanh\DMC_kha-nang-khach-hang-quay-lai\BaoCao_AI_DuDoanKhachHangQuayLai.md"
    docx = r"c:\Users\Admin\ducmanh\DMC_kha-nang-khach-hang-quay-lai\BaoCao_AI_DuDoanKhachHangQuayLai.docx"
    imgs = r"c:\Users\Admin\ducmanh\DMC_kha-nang-khach-hang-quay-lai\images"
    convert_md_to_docx(md, docx, imgs)
