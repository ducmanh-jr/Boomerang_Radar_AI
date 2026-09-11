import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(val='single', sz='12', color='000000')
    """
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for b_name in ['top', 'left', 'bottom', 'right']:
        if b_name in kwargs:
            val = kwargs[b_name].get('val', 'single')
            sz = kwargs[b_name].get('sz', '4')
            color = kwargs[b_name].get('color', 'CCCCCC')
            node = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(node)
        else:
            node = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(node)
    tcPr.append(tcBorders)

def add_header_footer(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    # Header for page 2+
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Customer Repurchase Prediction AI System  |  Báo Cáo Kỹ Thuật AI")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(9)
    hrun.font.color.rgb = RGBColor(0, 0, 0)
    
    # Simple thin bottom border for header
    hpBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="000000"/></w:pBdr>')
    hp._element.get_or_add_pPr().append(hpBdr)

    # Footer for page 2+
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    frun1 = fp.add_run("Trang ")
    frun1.font.name = "Times New Roman"
    frun1.font.size = Pt(9)
    frun1.font.color.rgb = RGBColor(0, 0, 0)
    
    fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    fp._element.append(fldSimple)

def build_academic_docx():
    md_path = r"c:\Users\Admin\ducmanh\DMC_kha-nang-khach-hang-quay-lai\BaoCao_AI_DuDoanKhachHangQuayLai.md"
    docx_path = r"c:\Users\Admin\ducmanh\DMC_kha-nang-khach-hang-quay-lai\BaoCao_AI_DuDoanKhachHangQuayLai.docx"
    
    doc = Document()
    
    # 1 inch page margins everywhere
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    add_header_footer(doc)

    # ==================== TRANG BÌA ĐƠN GIẢN ĐẸP (CLEAN ACADEMIC COVER PAGE) ====================
    doc.add_paragraph().paragraph_format.space_before = Pt(36)
    
    # Institution Header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(24)
    r_inst = p_inst.add_run("ENTERPRISE AI IMPLEMENTATION & ANALYTICS DIVISION\n***")
    r_inst.bold = True
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(12)
    r_inst.font.color.rgb = RGBColor(0, 0, 0)

    # Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("BÁO CÁO TỔNG QUAN HỆ THỐNG AI DỰ ĐOÁN KHẢ NĂNG KHÁCH HÀNG QUAY LẠI MUA HÀNG")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(36)
    r_sub = p_sub.add_run("CUSTOMER REPURCHASE PREDICTION AI SYSTEM\nEnd-to-End Machine Learning Pipeline & Enterprise Architecture")
    r_sub.italic = True
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(0, 0, 0)

    # Divider Line
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(72)
    r_line = p_line.add_run("—" * 35)
    r_line.font.name = "Times New Roman"
    r_line.font.size = Pt(11)
    r_line.font.color.rgb = RGBColor(0, 0, 0)

    # Author Metadata Box (Clean Black on White)
    meta_items = [
        ("Tác giả báo cáo:", "Principal ML System Engineer & Head of Data Science"),
        ("Đơn vị thực hiện:", "Enterprise AI Implementation & Analytics Division"),
        ("Ngày phát hành:", "07/09/2026"),
        ("Trạng thái hệ thống:", "Production-Ready / Enterprise Scalable Infrastructure")
    ]
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (k, v) in enumerate(meta_items):
        c1 = meta_table.cell(idx, 0)
        c2 = meta_table.cell(idx, 1)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(0, 0, 0)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()

    # ==================== TRANG MỤC LỤC & DANH MỤC (TOC & FIGURES) ====================
    # Title TOC
    p_toc_h = doc.add_paragraph()
    p_toc_h.paragraph_format.space_before = Pt(12)
    p_toc_h.paragraph_format.space_after = Pt(12)
    r_toc_h = p_toc_h.add_run("MỤC LỤC TỔNG QUAN")
    r_toc_h.bold = True
    r_toc_h.font.name = "Times New Roman"
    r_toc_h.font.size = Pt(16)
    r_toc_h.font.color.rgb = RGBColor(0, 0, 0)

    toc_items = [
        ("PHẦN I: ĐÁNH GIÁ VÀ KIỂM TOÁN TÀI LIỆU HỆ THỐNG CŨ (AUDIT REPORT)", "Trang 3"),
        ("  1.1. Nhận xét & Hạn chế của Hệ thống Cũ khi Triển khai Thực tế", "Trang 3"),
        ("  1.2. Các Điểm mạnh Cần Tiếp tục Phát huy", "Trang 3"),
        ("PHẦN II: CƠ SỞ TOÁN HỌC VÀ LÝ THUYẾT HỌC MÁY CHUYÊN SÂU", "Trang 4"),
        ("  2.1. Hồi quy Logistic & Động lực học Tối ưu hóa Gradient Descent", "Trang 4"),
        ("  2.2. Cây Quyết định & Các Chỉ số Vùng Không Nguyên chất (Impurity)", "Trang 5"),
        ("  2.3. Rừng Ngẫu nhiên & Lý thuyết Ensemble Learning (Bagging Proof)", "Trang 6"),
        ("  2.4. Thuật toán XGBoost & Khai triển Taylor Bậc Hai", "Trang 7"),
        ("  2.5. Khái niệm Shapley Values trong Explainable AI (XAI)", "Trang 8"),
        ("  2.6. Mô hình Phân tích Sinh tồn Cox (Cox Proportional Hazards)", "Trang 8"),
        ("PHẦN III: KIẾN TRÚC HỆ THỐNG ML ENTERPRISE VÀ HẠ TẦNG MLOPS", "Trang 9"),
        ("  3.1. Các Thành phần Hạ tầng Hệ thống (Event Ingestion, Feature Store)", "Trang 9"),
        ("  3.2. Cân bằng Độ lệch - Phương sai trong Hệ thống Enterprise", "Trang 9"),
        ("  3.3. Chỉ số SLA Độ trễ và Payload Chuẩn API (REST Contract)", "Trang 10"),
        ("PHẦN IV: QUY TRÌNH XỬ LÝ DỮ LIỆU THỰC NGHIỆM VÀ ANTI-DATA LEAKAGE", "Trang 11"),
        ("  4.1. Từ điển Đặc trưng Hệ thống (System Feature Dictionary - 19 Features)", "Trang 11"),
        ("  4.2. Phân tích Phân phối Dữ liệu EDA", "Trang 12"),
        ("  4.3. Pipeline Kỹ thuật Đặc trưng & Biến đổi Dữ liệu", "Trang 12"),
        ("  4.4. Khung Xử lý Chống Rò rỉ Dữ liệu (Anti-Data Leakage Pipeline Frame)", "Trang 13"),
        ("  4.5. Chiến lược Kiểm định Chéo (K-Fold & Stratified Temporal Split)", "Trang 13"),
        ("PHẦN V: KẾT QUẢ THỰC NGHIỆM VÀ PHÂN TÍCH ĐÁNH GIÁ MÔ HÌNH", "Trang 14"),
        ("  5.1. Bảng So sánh Hiệu năng Mô hình trên Tập Kiểm thử (1,126 Mẫu)", "Trang 14"),
        ("  5.2. Biểu đồ Cột So sánh Chi tiết các Chỉ số", "Trang 14"),
        ("  5.3. Ma trận Nhầm lẫn (Confusion Matrices)", "Trang 15"),
        ("  5.4. Đồ thị Đường cong ROC và Precision-Recall", "Trang 15"),
        ("  5.5. Phân tích Chi tiết Mô hình Quán quân (Random Forest Breakdown)", "Trang 16"),
        ("  5.6. Phân tích Độ phức tạp Mô hình (Overfitting vs Underfitting)", "Trang 16"),
        ("  5.7. Tối ưu hóa Ngưỡng Phân quyết Nghiệp vụ (Threshold Tuning)", "Trang 17"),
        ("  5.8. Khung Quyết định Lựa chọn Mô hình Enterprise", "Trang 17"),
        ("  5.9. Kiểm định Ý nghĩa Thống kê (Statistical Significance Testing)", "Trang 18"),
        ("PHẦN VI: GIẢI THÍCH MÔ HÌNH CHUYÊN SÂU BẰNG SHAP (MODEL INTERPRETABILITY)", "Trang 19"),
        ("  6.1. Tầm quan trọng Đặc trưng Toàn cục (Global Feature Importances)", "Trang 19"),
        ("  6.2. Phân tích SHAP Waterfall Cá thể hóa", "Trang 19"),
        ("  6.3. Top 5 Đặc trưng Quyết định (Core Driving Features)", "Trang 20"),
        ("PHẦN VII: CRM INTEGRATION PLAYBOOK VÀ QUY TRÌNH MLOPS CI/CD", "Trang 21"),
        ("  7.1. Chiến lược Phân hạng Khách hàng và CRM Automation Playbook", "Trang 21"),
        ("  7.2. Quy trình MLOps CI/CD & Giám sát Trôi Dữ liệu (Drift Monitoring)", "Trang 22"),
        ("PHẦN VIII: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN HỆ THỐNG", "Trang 23"),
        ("  8.1. Tóm tắt Đóng góp và Giá trị Kinh tế (Business Value Summary)", "Trang 23"),
        ("  8.2. Kế hoạch Phát triển Tương lai (Future Roadmap)", "Trang 23")
    ]

    toc_table = doc.add_table(rows=len(toc_items), cols=2)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (title, page_str) in enumerate(toc_items):
        c_t = toc_table.cell(idx, 0)
        c_p = toc_table.cell(idx, 1)
        set_cell_margins(c_t, top=30, bottom=30, left=40, right=40)
        set_cell_margins(c_p, top=30, bottom=30, left=40, right=40)
        
        is_main = title.startswith("PHẦN")
        pt = c_t.paragraphs[0]
        rt = pt.add_run(title)
        rt.bold = is_main
        rt.font.name = "Times New Roman"
        rt.font.size = Pt(10.5 if is_main else 10)
        rt.font.color.rgb = RGBColor(0, 0, 0)
        
        pp = c_p.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp = pp.add_run(page_str)
        rp.bold = is_main
        rp.font.name = "Times New Roman"
        rp.font.size = Pt(10.5 if is_main else 10)
        rp.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph().paragraph_format.space_before = Pt(18)

    # DANH MỤC HÌNH VẼ (LIST OF FIGURES)
    p_fig_h = doc.add_paragraph()
    p_fig_h.paragraph_format.space_before = Pt(12)
    p_fig_h.paragraph_format.space_after = Pt(12)
    r_fig_h = p_fig_h.add_run("DANH MỤC HÌNH VẼ (24 HÌNH MINH HỌA)")
    r_fig_h.bold = True
    r_fig_h.font.name = "Times New Roman"
    r_fig_h.font.size = Pt(14)
    r_fig_h.font.color.rgb = RGBColor(0, 0, 0)

    figures_list = [
        ("Hình 1", "Kiến trúc tổng thể Hệ thống AI Dự đoán Khách hàng Quay lại (Enterprise Architecture)"),
        ("Hình 2", "Phân tích sự cân bằng giữa Độ lệch (Bias), Phương sai (Variance) & Ensemble Learning"),
        ("Hình 3", "Đồ thị hàm Sigmoid σ(z) và tính chất đạo hàm đạt cực đại σ'(0) = 0.25"),
        ("Hình 4", "So sánh hình học giữa Shannon Entropy H(p) và Gini Impurity G(p)"),
        ("Hình 5", "Chứng minh tính chất giảm Phương sai (Variance Reduction) của Bagging"),
        ("Hình 6", "Mặt phẳng tổn thất 3D biểu diễn tính lồi (Convexity) của hàm Log-Loss"),
        ("Hình 7", "Quy trình nghiêm ngặt phân tách dữ liệu chống rò rỉ (Anti-Data Leakage Pipeline)"),
        ("Hình 8", "Phân tích phân phối tần suất của các đặc trưng cốt lõi (Tenure, Satisfaction, Recency, Cashback)"),
        ("Hình 9", "Biểu đồ cột so sánh chi tiết các chỉ số Accuracy, Precision, Recall, F1, ROC-AUC"),
        ("Hình 10", "Ma trận nhầm lẫn (Confusion Matrices) đánh giá TP, FP, TN, FN trên tập Test"),
        ("Hình 11", "Đồ thị đường cong ROC và Precision-Recall Curve (Random Forest AUC = 0.9241)"),
        ("Hình 12", "Xếp hạng tầm quan trọng của 10 đặc trưng hàng đầu theo MDI và SHAP Value"),
        ("Hình 13", "Đồ thị SHAP Waterfall giải thích chi tiết các yếu tố tăng/giảm xác suất cá thể"),
        ("Hình 14", "Sơ đồ phân hạng rủi ro và Kịch bản hành động tự động hóa CRM Marketing"),
        ("Hình 15", "Quy trình MLOps khép kín từ CI/CD Deployment đến Tự động Re-training khi Drift"),
        ("Hình 16", "Không gian mặt phẳng tổn thất và quỹ đạo hội tụ của SGD, Batch GD và Adam Optimizer"),
        ("Hình 17", "Phân biệt giữa K-Fold Cross Validation tiêu chuẩn và Stratified Temporal Split"),
        ("Hình 18", "Biểu diễn hình học không gian ràng buộc của Regularization L1 (Lasso) và L2 (Ridge)"),
        ("Hình 19", "So sánh cơ chế hoạt động giữa Bagging (Parallel Voting) và Boosting (Sequential Error)"),
        ("Hình 20", "Đồ thị độ phức tạp mô hình (Model Capacity) thể hiện điểm cân bằng Overfitting/Underfitting"),
        ("Hình 21", "Phân tích sự đánh đổi Precision-Recall theo ngưỡng phân quyết và tối ưu hóa lợi nhuận"),
        ("Hình 22", "Quy trình biến đổi đặc trưng từ Raw Data sang Feature Vector phục vụ mô hình ML"),
        ("Hình 23", "Cây quyết định lựa chọn mô hình dựa trên tiêu chí Accuracy, Latency SLA và Interpretability"),
        ("Hình 24", "Kết quả kiểm định thống kê Paired t-test và McNemar's Test (p < 0.001)")
    ]

    fig_table = doc.add_table(rows=len(figures_list), cols=2)
    fig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (f_num, f_desc) in enumerate(figures_list):
        c_n = fig_table.cell(idx, 0)
        c_d = fig_table.cell(idx, 1)
        set_cell_margins(c_n, top=20, bottom=20, left=40, right=40)
        set_cell_margins(c_d, top=20, bottom=20, left=40, right=40)
        
        pn = c_n.paragraphs[0]
        rn = pn.add_run(f_num)
        rn.bold = True
        rn.font.name = "Times New Roman"
        rn.font.size = Pt(9.5)
        rn.font.color.rgb = RGBColor(0, 0, 0)
        
        pd = c_d.paragraphs[0]
        rd = pd.add_run(f_desc)
        rd.font.name = "Times New Roman"
        rd.font.size = Pt(9.5)
        rd.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()

    # ==================== PARSING MARKDOWN CONTENT (PURE BLACK ON WHITE) ====================
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_data = []
    in_code = False
    code_lines = []

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Code block (Clean Black Text on White background with subtle thin border)
        if line.startswith('```'):
            if in_code:
                code_text = '\n'.join(code_lines)
                
                c_tbl = doc.add_table(rows=1, cols=1)
                c_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                c_cell = c_tbl.cell(0, 0)
                set_cell_margins(c_cell, top=120, bottom=120, left=160, right=160)
                set_cell_border(c_cell, top=dict(sz='4', color='CCCCCC'),
                                        bottom=dict(sz='4', color='CCCCCC'),
                                        left=dict(sz='4', color='CCCCCC'),
                                        right=dict(sz='4', color='CCCCCC'))
                
                cp = c_cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                cp.paragraph_format.line_spacing = 1.15
                c_run = cp.add_run(code_text)
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(9.5)
                c_run.font.color.rgb = RGBColor(0, 0, 0)
                
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Markdown Tables (Pure Black Text on White Background)
        if '|' in line and not line.startswith('!['):
            parts = [cell.strip() for cell in line.split('|')[1:-1]]
            if len(parts) > 0:
                if all(re.match(r'^:?-+:?$', p) for p in parts):
                    i += 1
                    continue
                in_table = True
                table_data.append(parts)
                i += 1
                continue
        else:
            if in_table and len(table_data) > 0:
                cols = max(len(row) for row in table_data)
                rows = len(table_data)
                tbl = doc.add_table(rows=rows, cols=cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                for r_idx, row_vals in enumerate(table_data):
                    for c_idx, val in enumerate(row_vals):
                        if c_idx < cols:
                            cell = tbl.cell(r_idx, c_idx)
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                            
                            # Thin black/gray border for table
                            b_bottom = dict(sz='8' if r_idx == 0 else '4', color='000000' if r_idx == 0 else 'CCCCCC')
                            set_cell_border(cell, top=dict(sz='4', color='CCCCCC'),
                                                    bottom=b_bottom,
                                                    left=dict(sz='4', color='CCCCCC'),
                                                    right=dict(sz='4', color='CCCCCC'))
                            
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            
                            run = p.add_run(val)
                            run.bold = (r_idx == 0)
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10)
                            run.font.color.rgb = RGBColor(0, 0, 0)
                                
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                table_data = []
                in_table = False

        if not line.strip():
            i += 1
            continue

        # Horizontal Rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/></w:pBdr>')
            p._element.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Images (6 inches wide, centered)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line.strip())
        if img_match:
            alt_text, img_rel_path = img_match.groups()
            img_abs_path = os.path.normpath(os.path.join(r"c:\Users\Admin\ducmanh\DMC_kha-nang-khach-hang-quay-lai", img_rel_path))
            if os.path.exists(img_abs_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(14)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(img_abs_path, width=Inches(6.0))
                
                if i + 1 < len(lines) and lines[i+1].strip().startswith('*Hình'):
                    caption_line = lines[i+1].strip()
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(0)
                    p_cap.paragraph_format.space_after = Pt(14)
                    run_cap = p_cap.add_run(caption_line.strip('*'))
                    run_cap.italic = True
                    run_cap.font.name = "Times New Roman"
                    run_cap.font.size = Pt(9.5)
                    run_cap.font.color.rgb = RGBColor(0, 0, 0)
                    i += 1
            i += 1
            continue

        # Headings (Pure Black on White)
        if line.startswith('# '):
            p = doc.add_paragraph(style='Heading 1')
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:])
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith('## '):
            p = doc.add_paragraph(style='Heading 2')
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[3:])
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith('### '):
            p = doc.add_paragraph(style='Heading 3')
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[4:])
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[5:])
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line[2:])
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif re.match(r'^\d+\.\s', line):
            content = re.sub(r'^\d+\.\s', '', line)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(content)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)

        i += 1

    doc.save(docx_path)
    print(f"Successfully built academic black-on-white DOCX at: {docx_path}")

if __name__ == '__main__':
    build_academic_docx()
